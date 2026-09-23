#!/usr/bin/env python3
"""
EV1O -- OXPHOS-vs-pathways correlation analysis (backs the manuscript's "visibly
correlated with increased FAO and BCAA, and reduced 1C metabolism" statement).
Single self-contained script, no intermediate cache files.

Inputs:

EV tables:
      EV_Table_8*.xlsx   -- MRC5 proteomics
      EV_Table_10*.xlsx  -- 30-model TIS meta-analysis (Anerillas)
      EV_Table_12*.xlsx  -- Payea et al. 2024 IMR90 etoposide dataset

Pathway gene lists (OXPHOS subunits, fatty acid oxidation, BCAA, NEAA, sulfur
metabolism, methionine metabolism, 1C/folate) and MITO_GE_GENES (gene symbols
annotated to the MitoCarta3.0 "Translation / mtDNA maintenance / mtRNA metabolism"
pathways; loaded but not one of the four categories actually plotted/correlated in
the final EV1O figure) are fixed reference sets, hardcoded directly below rather
than read from a separate file.

Output:
  Correlation_OXPHOS_vs_pathways.png/svg
  Correlation_analysis.xlsx
"""
import glob
import os
import numpy as np
import pandas as pd
from scipy import stats
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.getcwd()
print('Running from:', HERE)
DATA_DIR = 'data'


def find_file(patterns, label):
    """Find a required input by glob pattern(s): checks ./data/ first, then
    the current folder, and raises FileNotFoundError naming the missing file."""
    for pat in patterns:
        matches = glob.glob(f'{DATA_DIR}/{pat}') or glob.glob(pat)
        if matches:
            return matches[0]
    raise FileNotFoundError(
        f"Missing required input for {label}: none of {patterns} found in "
        f"./{DATA_DIR}/ or next to this script. Check the filename/location and rerun."
    )


# ---- EV tables ----
EV_TABLE_8_PATH = find_file(['EV_Table_8*.xlsx'], 'EV Table 8')
ANERILLAS_PATH = find_file(['EV_Table_10*.xlsx'], '30-model TIS meta-analysis (Anerillas)')
PAYEA_PATH = find_file(
    ['EV_Table_12*.xlsx'],
    'EV Table 12 / Payea et al. 2024 IMR90 etoposide dataset')

print('All required inputs found:')
for _label, _path in [('EV_TABLE_8_PATH', EV_TABLE_8_PATH), ('ANERILLAS_PATH', ANERILLAS_PATH),
                       ('PAYEA_PATH', PAYEA_PATH)]:
    print(f'  {_label}: {_path}')
print()

# Hardcoded (curated, cross-checked against MitoCarta3.0; see
# no external file needed.
MITO_GE_GENES = {
    'AARS2', 'ALKBH1', 'ANGEL2', 'APEX1', 'ATAD3A', 'ATAD3B', 'AURKAIP1', 'CARS2', 'CDK5RAP1',
    'CHCHD1', 'COA3', 'COX14', 'DAP3', 'DARS2', 'DDX28', 'DHX30', 'DNA2', 'DUS2', 'EARS2',
    'ELAC2', 'ENDOG', 'ERAL1', 'EXD2', 'EXOG', 'FARS2', 'FASTK', 'FASTKD1', 'FASTKD2',
    'FASTKD3', 'FASTKD5', 'GADD45GIP1', 'GARS1', 'GATB', 'GATC', 'GFM1', 'GFM2', 'GRSF1',
    'GTPBP10', 'GTPBP3', 'GUF1', 'HARS2', 'HEMK1', 'HSD17B10', 'IARS2', 'KARS1', 'KGD4',
    'LACTB2', 'LARS2', 'LIG3', 'LRPPRC', 'MALSU1', 'MARS2', 'METAP1D', 'METTL15', 'METTL17',
    'METTL5', 'METTL8', 'MGME1', 'MIEF1', 'MPV17L2', 'MRM1', 'MRM2', 'MRM3', 'MRPL1', 'MRPL10',
    'MRPL11', 'MRPL12', 'MRPL13', 'MRPL14', 'MRPL15', 'MRPL16', 'MRPL17', 'MRPL18', 'MRPL19',
    'MRPL2', 'MRPL20', 'MRPL21', 'MRPL22', 'MRPL23', 'MRPL24', 'MRPL27', 'MRPL28', 'MRPL3',
    'MRPL30', 'MRPL32', 'MRPL33', 'MRPL34', 'MRPL35', 'MRPL36', 'MRPL37', 'MRPL38', 'MRPL39',
    'MRPL4', 'MRPL40', 'MRPL41', 'MRPL42', 'MRPL43', 'MRPL44', 'MRPL46', 'MRPL47', 'MRPL48',
    'MRPL49', 'MRPL50', 'MRPL51', 'MRPL52', 'MRPL53', 'MRPL54', 'MRPL55', 'MRPL57', 'MRPL58',
    'MRPL9', 'MRPS10', 'MRPS11', 'MRPS12', 'MRPS14', 'MRPS15', 'MRPS16', 'MRPS17', 'MRPS18A',
    'MRPS18B', 'MRPS18C', 'MRPS2', 'MRPS21', 'MRPS22', 'MRPS23', 'MRPS24', 'MRPS25', 'MRPS26',
    'MRPS27', 'MRPS28', 'MRPS30', 'MRPS31', 'MRPS33', 'MRPS34', 'MRPS35', 'MRPS5', 'MRPS6',
    'MRPS7', 'MRPS9', 'MRRF', 'MTERF3', 'MTERF4', 'MTFMT', 'MTG1', 'MTG2', 'MTIF2', 'MTIF3',
    'MTO1', 'MTPAP', 'MTRES1', 'MTRF1', 'MTRF1L', 'MUTYH', 'NARS2', 'NGRN', 'NOA1', 'NSUN2',
    'NSUN4', 'OGG1', 'OSGEPL1', 'OXA1L', 'PARS2', 'PDE12', 'PDF', 'PIF1', 'PNPT1', 'POLB',
    'POLDIP2', 'POLG', 'POLG2', 'POLQ', 'POLRMT', 'PPA2', 'PRORP', 'PTCD1', 'PTCD2', 'PTCD3',
    'PUS1', 'PUSL1', 'QRSL1', 'QTRT1', 'RARS2', 'RBFA', 'RCC1L', 'RECQL4', 'REXO2', 'RMND1',
    'RNASEH1', 'RPUSD3', 'RPUSD4', 'SARS2', 'SLIRP', 'SSBP1', 'SUPV3L1', 'TACO1', 'TARS2',
    'TBRG4', 'TEFM', 'TFAM', 'TFB1M', 'TFB2M', 'THG1L', 'TIMM21', 'TOP3A', 'TRIT1', 'TRMT1',
    'TRMT10C', 'TRMT2B', 'TRMT5', 'TRMT61B', 'TRMU', 'TRNT1', 'TRUB2', 'TSFM', 'TUFM', 'TWNK',
    'UNG', 'VARS2', 'WARS2', 'YARS2', 'YBEY', 'YRDC'
}  # n=222

# ============================================================== 1. build gene-level data
RENAME = {
    'ATP5A1': 'ATP5F1A', 'ATP5B': 'ATP5F1B', 'ATP5C1': 'ATP5F1C', 'ATP5D': 'ATP5F1D',
    'ATP5F1': 'ATP5PB', 'ATP5H': 'ATP5PD', 'ATP5I': 'ATP5ME', 'ATP5J': 'ATP5PF',
    'ATP5J2': 'ATP5MF', 'ATP5L': 'ATP5MG', 'ATP5O': 'ATP5PO',
    'UQCRFS1;UQCRFS1P1': 'UQCRFS1',
    'SQRDL': 'SQOR',
}


def rn(g):
    return RENAME.get(g, g)


records = []
model_background = {}

# Pathway gene lists (fixed reference sets, not derived from any input file).
gene_lists = {
    'OXPHOS_subunits': [
        'ATP5F1A', 'ATP5F1B', 'ATP5F1C', 'ATP5F1D', 'ATP5PB', 'ATP5PD', 'ATP5ME', 'ATP5PF',
        'ATP5MF', 'ATP5MG', 'ATP5PO', 'COX4I1', 'COX5A', 'COX5B', 'COX6A1', 'COX6B1', 'COX6C',
        'COX7A2', 'COX7A2L', 'COX7C', 'CYC1', 'CYCS', 'HCCS', 'MT-ATP6', 'MT-ATP8', 'MT-CO1',
        'MT-CO2', 'MT-ND1', 'NDUFA10', 'NDUFA11', 'NDUFA12', 'NDUFA13', 'NDUFA2', 'NDUFA3',
        'NDUFA4', 'NDUFA5', 'NDUFA6', 'NDUFA7', 'NDUFA8', 'NDUFA9', 'NDUFAB1', 'NDUFB1',
        'NDUFB10', 'NDUFB11', 'NDUFB2', 'NDUFB3', 'NDUFB4', 'NDUFB5', 'NDUFB6', 'NDUFB7',
        'NDUFB8', 'NDUFB9', 'NDUFS1', 'NDUFS2', 'NDUFS3', 'NDUFS4', 'NDUFS5', 'NDUFS7',
        'NDUFS8', 'NDUFV1', 'NDUFV2', 'SDHA', 'SDHB', 'UQCR11', 'UQCRB', 'UQCRC1', 'UQCRC2',
        'UQCRFS1', 'UQCRH', 'UQCRQ',
    ],
    'Fatty_acid_oxidation': [
        'ACAA1', 'ACAA2', 'ACACA', 'ACAD10', 'ACAD11', 'ACADM', 'ACADSB', 'ACADVL', 'ACAT1',
        'ACOT13', 'ACOT7', 'ACOT9', 'ACSF2', 'ACSF3', 'ACSL1', 'AGK', 'CPT1A', 'CPT2', 'CRAT',
        'CROT', 'CYB5R3', 'DBI', 'DECR1', 'DHRS1', 'ECH1', 'ECHDC1', 'ECHS1', 'ECI1', 'ECI2',
        'ETFA', 'ETFB', 'ETFDH', 'FASN', 'FDPS', 'FDXR', 'GCSH', 'HADH', 'HADHA', 'HADHB',
        'HINT2', 'HSD17B10', 'HSD17B4', 'IDI1', 'LACTB', 'LYPLA1', 'MCEE', 'MGST3', 'MUT',
        'NDUFAB1', 'OSBPL1A', 'PCCB', 'PLSCR3', 'PRDX6', 'PTGES2', 'PTPMT1', 'SCP2', 'SLC25A1',
        'SLC25A20', 'SPTLC2', 'TAMM41', 'TSPO',
    ],
    'BCAA_metabolism': [
        'ACADSB', 'ACAT1', 'ALDH6A1', 'BCAT2', 'BCKDHA', 'DBT', 'DLD', 'ECHS1', 'ETFA', 'ETFB',
        'ETFDH', 'HADHA', 'HIBADH', 'HIBCH', 'HMGCL', 'HSD17B10', 'IVD', 'MCCC1', 'MCCC2',
    ],
    'NEAA_metabolism': [
        'ALDH18A1', 'ASL', 'ASNS', 'ASS1', 'BCAT2', 'GLS', 'GOT1', 'GOT2', 'PHGDH', 'PSAT1',
        'PSPH', 'PYCR1', 'PYCR2', 'SHMT2',
    ],
    'Sulfur_metabolism': ['ETHE1', 'GOT2', 'MPST', 'MSRA', 'SQOR', 'TST'],
    'Methionine_metabolism_GOBP': [
        'ADI1', 'AHCY', 'AHCYL1', 'AHCYL2', 'APIP', 'ENOPH1', 'MAT2A', 'MAT2B', 'MRI1', 'MSRA',
        'MTAP', 'SMS',
    ],
    'OneC_folate_metabolism': [
        'ALDH1L1', 'MTHFD1', 'MTHFR', 'SHMT1', 'DHFR', 'TYMS', 'MTHFD1L', 'MTHFD2', 'ALDH1L2',
        'SHMT2', 'GART', 'ATIC',
    ],
}
print('Canonical gene-list sizes:')
for k, v in gene_lists.items():
    print(' ', k, len(v))

# sheet_name=0 (by position, not by name) and a column-whitespace strip: prior EV
# tables in this project have shown up with renamed sheets and/or stray leading
# spaces on column names in different exported copies of the "same" file.
cam = pd.read_excel(EV_TABLE_8_PATH, sheet_name=0)
cam.columns = cam.columns.str.strip()
cam = cam.dropna(subset=['Gene names']).drop_duplicates(subset=['Gene names'], keep='first')
cam['Gene names'] = cam['Gene names'].map(rn)
cam = cam.set_index('Gene names')
own_fc = cam['Log2FC Sen CTRL vs Prolif CTRL']
own_bg_mean = own_fc.dropna().mean()

MRC5_LABEL = 'MRC5'
model_background[MRC5_LABEL] = (own_fc.dropna() - own_bg_mean).values
print('\nMRC5 full-proteome background size:', len(model_background[MRC5_LABEL]))
for cat, genes in gene_lists.items():
    for g in genes:
        if g in own_fc.index and pd.notna(own_fc[g]):
            records.append(dict(model=MRC5_LABEL, category=cat, gene=g,
                                 log2FC_rel=own_fc[g] - own_bg_mean))

ge_genes_mrc5 = [g for g in own_fc.dropna().index if g in MITO_GE_GENES]
for g in ge_genes_mrc5:
    records.append(dict(model=MRC5_LABEL, category='Mito_gene_expression',
                         gene=g, log2FC_rel=own_fc[g] - own_bg_mean))
print('MRC5 Mito_gene_expression genes:', len(ge_genes_mrc5), '(not used by the final EV1O plot either way.)')

payea = pd.read_excel(PAYEA_PATH, sheet_name=0)
payea.columns = payea.columns.str.strip()
payea['Symbol'] = payea['Symbol'].map(rn)
cyc_cols = ['Cyc_1', 'Cyc_2', 'Cyc_3']
etop_cols = ['Etop_1', 'Etop_2', 'Etop_3']
cyc_mean = payea[cyc_cols].mean(axis=1)
etop_mean = payea[etop_cols].mean(axis=1)
valid = (cyc_mean > 0) & (etop_mean > 0)
payea_fc_raw = pd.Series(np.log2(etop_mean[valid] / cyc_mean[valid]).values, index=payea.loc[valid, 'Symbol'].values)
payea_fc_raw = payea_fc_raw[~payea_fc_raw.index.duplicated()]
payea_bg_mean = payea_fc_raw.mean()

PAYEA_LABEL = 'IMR90 (etoposide, Payea 2024)'
model_background[PAYEA_LABEL] = (payea_fc_raw.dropna() - payea_bg_mean).values
print('Payea full-proteome background size:', len(model_background[PAYEA_LABEL]))
for cat, genes in gene_lists.items():
    for g in genes:
        if g in payea_fc_raw.index and pd.notna(payea_fc_raw[g]):
            records.append(dict(model=PAYEA_LABEL, category=cat, gene=g,
                                 log2FC_rel=payea_fc_raw[g] - payea_bg_mean))

ge_genes = [g for g in payea_fc_raw.index if g in MITO_GE_GENES]
for g in ge_genes:
    records.append(dict(model=PAYEA_LABEL, category='Mito_gene_expression', gene=g,
                         log2FC_rel=payea_fc_raw[g] - payea_bg_mean))
print('Payea Mito_gene_expression genes:', len(ge_genes))

CELLTYPE_SHEETS = {'WI38': 'WI38', 'BJ': 'BJ', 'HSAEC': 'HSAEC', 'HEKn': 'HEKn', 'HCAEC': 'HCAEC',
                    'HUVEC': 'HUVEC', 'BMMSC': 'BMMSC', 'HVSMC': 'HVSMC', 'HSKM': 'HSKM', 'PBMC': 'PBMC',
                    'PreAdipo': 'PreAdipo', 'NHO': 'NHO', 'NHA': 'NHA', 'HEMn': 'HEM'}

for ct, sheet in CELLTYPE_SHEETS.items():
    d = pd.read_excel(ANERILLAS_PATH, sheet_name=sheet)
    d.columns = d.columns.str.strip()
    d = d.dropna(subset=['Gene Symbol']).drop_duplicates(subset=['Gene Symbol'], keep='first')
    d['Gene Symbol'] = d['Gene Symbol'].map(rn)
    d = d.set_index('Gene Symbol')
    for cond in ['CTIS', 'IRIS']:
        diff_col = f"Student's T-test Difference {ct}_{cond}_{ct}_P"
        if diff_col not in d.columns:
            continue
        diff = d[diff_col]
        bg_mean = diff.mean()
        model_label = f'{ct} {cond}'
        model_background[model_label] = (diff.dropna() - bg_mean).values
        for cat, genes in gene_lists.items():
            for g in genes:
                if g in diff.index and pd.notna(diff[g]):
                    records.append(dict(model=model_label, category=cat, gene=g,
                                         log2FC_rel=diff[g] - bg_mean))
        ge_genes_ct = [g for g in diff.index if g in MITO_GE_GENES]
        for g in ge_genes_ct:
            if pd.notna(diff[g]):
                records.append(dict(model=model_label, category='Mito_gene_expression', gene=g,
                                     log2FC_rel=diff[g] - bg_mean))

gene_df = pd.DataFrame.from_records(records)
print('\ntotal gene-level records:', len(gene_df))
print('models built:', sorted(gene_df.model.unique()))
assert MRC5_LABEL in gene_df.model.unique(), f'MRC5 missing from gene_df -- check {EV_TABLE_8_PATH}'

# ============================================================== 2. competitive permutation test
N_PERM = 2000
rng = np.random.default_rng(0)

perm_rows = []
for (model, cat), sub in gene_df.groupby(['model', 'category']):
    vals = sub['log2FC_rel'].values
    n = len(vals)
    obs_mean = vals.mean()
    bg = model_background[model]
    B = len(bg)
    if n == 0 or n > B:
        perm_rows.append(dict(model=model, category=cat, n=n, obs_mean=np.nan, perm_p=np.nan))
        continue
    keys = rng.random((N_PERM, B))
    idx = np.argpartition(keys, n, axis=1)[:, :n]
    null_means = bg[idx].mean(axis=1)
    p = (np.sum(np.abs(null_means) >= abs(obs_mean)) + 1) / (N_PERM + 1)
    perm_rows.append(dict(model=model, category=cat, n=n, obs_mean=obs_mean, perm_p=p))
perm = pd.DataFrame(perm_rows)
print('\npermutation test done:', perm.shape)
print(perm[perm.model == MRC5_LABEL].sort_values('category').to_string(index=False))

# ============================================================== 3. correlation analysis + figure
REF_MODELS = [MRC5_LABEL, PAYEA_LABEL]


def group_of(model):
    if model in REF_MODELS:
        return 'Reference'
    if model.endswith(' CTIS'):
        return 'CTIS'
    if model.endswith(' IRIS'):
        return 'IRIS'
    return 'Other'


CATS = {
    'OXPHOS_subunits': 'OXPHOS subunits',
    'Fatty_acid_oxidation': 'Fatty acid oxidation',
    'BCAA_metabolism': 'BCAA metabolism',
    'OneC_folate_metabolism': '1C metabolism (folate)',
}
wide = perm.pivot(index='model', columns='category', values='obs_mean')[list(CATS.keys())]
wide = wide.dropna()
wide['group'] = [group_of(m) for m in wide.index]
wide = wide.reset_index().rename(columns={'index': 'model'})
print('\nmodels with all 4 categories present:', len(wide))

pairs = [
    ('Fatty_acid_oxidation', 'OXPHOS vs Fatty acid oxidation'),
    ('BCAA_metabolism', 'OXPHOS vs BCAA metabolism'),
    ('OneC_folate_metabolism', 'OXPHOS vs 1C metabolism (folate)'),
]

stat_rows = []
for cat, label in pairs:
    r, p = stats.pearsonr(wide['OXPHOS_subunits'], wide[cat])
    rho, ps = stats.spearmanr(wide['OXPHOS_subunits'], wide[cat])
    stat_rows.append(dict(comparison=label, n=len(wide), pearson_r=r, pearson_p=p,
                           spearman_rho=rho, spearman_p=ps))
stats_df = pd.DataFrame(stat_rows)
print(stats_df.to_string(index=False))

# ---------------------------------------------------------------- design-system colors
INK, MUTED, GRID, SURFACE = '#0b0b0b', '#52514e', '#e1e0d9', '#fcfcfb'
GROUP_COLOR = {'Reference': '#2a78d6', 'CTIS': '#eb6834', 'IRIS': '#1baf7a'}
GROUP_ORDER = ['Reference', 'CTIS', 'IRIS']

fig, axes = plt.subplots(1, 3, figsize=(13.5, 4.6))

for ax, (cat, label) in zip(axes, pairs):
    x = wide['OXPHOS_subunits'].values
    y = wide[cat].values
    for g in GROUP_ORDER:
        sub = wide[wide.group == g]
        ax.scatter(sub['OXPHOS_subunits'], sub[cat], s=64, color=GROUP_COLOR[g],
                   edgecolor=SURFACE, linewidth=1.4, zorder=3, label=g)
    slope, intercept = np.polyfit(x, y, 1)
    xs = np.linspace(x.min(), x.max(), 50)
    ax.plot(xs, slope * xs + intercept, color=INK, lw=2, zorder=2, alpha=0.55)

    r_row = stats_df[stats_df.comparison.str.endswith(label.split('vs ')[1])].iloc[0]
    ax.text(0.03, 0.97, f"Pearson r = {r_row.pearson_r:.2f}\np = {r_row.pearson_p:.1e}",
            transform=ax.transAxes, fontsize=9.5, color=INK, va='top', ha='left',
            fontweight='bold')

    ax.axhline(0, color=GRID, lw=1.0, zorder=1)
    ax.axvline(0, color=GRID, lw=1.0, zorder=1)
    ax.set_xlabel('OXPHOS subunits\nmean log$_2$FC per model', fontsize=8.8, color=MUTED)
    ax.set_ylabel(f'{CATS[cat]}\nmean log$_2$FC per model', fontsize=8.8, color=MUTED)
    ax.set_title(label, fontsize=10.5, color=INK, fontweight='bold', pad=8)
    for s in ('top', 'right'):
        ax.spines[s].set_visible(False)
    ax.spines['left'].set_color(GRID)
    ax.spines['bottom'].set_color(GRID)
    ax.tick_params(colors=MUTED, labelsize=7.8)

axes[0].legend(loc='lower right', frameon=False, fontsize=8.5, handletextpad=0.4)

fig.suptitle('Per-model effect size correlates across pathway categories (n=30 TIS models)',
             fontsize=13, fontweight='bold', color=INK, x=0.015, ha='left', y=1.03)
fig.text(0.015, 0.965,
         'Each point = one model\'s permutation-test mean log$_2$FC for that category. Line = OLS fit.',
         fontsize=9, color=MUTED, ha='left', va='top')

fig.tight_layout(rect=[0, 0, 1, 0.90])
fig.savefig('Correlation_OXPHOS_vs_pathways.png', dpi=300, facecolor='white', bbox_inches='tight')
fig.savefig('Correlation_OXPHOS_vs_pathways.svg', facecolor='white', bbox_inches='tight')
print('\nsaved Correlation_OXPHOS_vs_pathways.png/svg')

# ================================================================== Excel export
FONT = 'Arial'
wb = Workbook()
wb.remove(wb.active)


def write_sheet(name, df, num_cols=()):
    ws = wb.create_sheet(name)
    ws.append(list(df.columns))
    for c in range(1, len(df.columns) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name=FONT, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='4A4A4A')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for _, r in df.iterrows():
        ws.append(list(r.values))
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(df.columns)):
        for cell in row:
            cell.font = Font(name=FONT, size=10)
    for i, col in enumerate(df.columns, start=1):
        letter = get_column_letter(i)
        if col in num_cols:
            for row in range(2, ws.max_row + 1):
                ws.cell(row=row, column=i).number_format = '0.000'
        maxlen = max([len(str(col))] + [len(str(v)) for v in df[col].values]) if len(df) else len(str(col))
        ws.column_dimensions[letter].width = min(max(maxlen + 2, 10), 42)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    return ws


display_df = wide[['model', 'group'] + list(CATS.keys())].copy()
display_df.columns = ['model', 'group'] + [CATS[c] for c in CATS.keys()]
write_sheet('Per_model_values', display_df, num_cols=list(CATS.values()))

stats_display = stats_df.rename(columns={
    'comparison': 'comparison', 'n': 'n_models', 'pearson_r': 'Pearson r', 'pearson_p': 'Pearson p',
    'spearman_rho': 'Spearman rho', 'spearman_p': 'Spearman p',
})
write_sheet('Correlation_stats', stats_display, num_cols=['Pearson r', 'Pearson p', 'Spearman rho', 'Spearman p'])

notes = wb.create_sheet('Read_me', 0)
notes_lines = [
    'Per_model_values: each of the 30 TIS models\' permutation-test mean log2FC '
    '(obs_mean, background-corrected) for OXPHOS subunits and the three pathways it is '
    'compared against. group = Reference (MRC5, IMR90) / CTIS (chemotherapy-induced, '
    'doxorubicin) / IRIS (irradiation-induced).',
    'Correlation_stats: Pearson r/p and Spearman rho/p for OXPHOS subunits vs each of '
    'Fatty acid oxidation, BCAA metabolism, and 1C metabolism (folate), across all 30 '
    'models. Backs up the manuscript\'s "visibly correlated" language in EV1O with '
    'an actual statistic.',
]
notes.column_dimensions['A'].width = 110
for i, line in enumerate(notes_lines, start=1):
    cell = notes.cell(row=i, column=1, value=line)
    cell.font = Font(name=FONT, size=11)
    cell.alignment = Alignment(wrap_text=True, vertical='top')
    notes.row_dimensions[i].height = 30

wb.save('Correlation_analysis.xlsx')
print('saved Correlation_analysis.xlsx to:', HERE)
